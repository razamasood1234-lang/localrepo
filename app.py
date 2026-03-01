from __future__ import annotations

from datetime import datetime
from pathlib import Path

from flask import Flask, redirect, render_template, request, url_for
from openpyxl import Workbook, load_workbook

app = Flask(__name__)

DATA_DIR = Path("data")
EXCEL_PATH = DATA_DIR / "submissions.xlsx"
HEADERS = ["Submitted At", "First Name", "Last Name", "Email", "Company"]


def ensure_workbook() -> None:
    """Create the Excel workbook with headers if it does not exist yet."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if EXCEL_PATH.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Form Submissions"
    sheet.append(HEADERS)

    for column in sheet.iter_cols(min_row=1, max_row=1):
        for cell in column:
            cell.font = cell.font.copy(bold=True)

    workbook.save(EXCEL_PATH)


@app.get("/")
def home() -> str:
    return render_template("index.html", success=False)


@app.post("/submit")
def submit():
    ensure_workbook()

    first_name = request.form.get("first_name", "").strip()
    last_name = request.form.get("last_name", "").strip()
    email = request.form.get("email", "").strip()
    company = request.form.get("company", "").strip()

    if not all([first_name, last_name, email, company]):
        return render_template(
            "index.html",
            success=False,
            error="Please fill in every field before submitting.",
            first_name=first_name,
            last_name=last_name,
            email=email,
            company=company,
        )

    workbook = load_workbook(EXCEL_PATH)
    sheet = workbook.active
    sheet.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        first_name,
        last_name,
        email,
        company,
    ])
    workbook.save(EXCEL_PATH)

    return redirect(url_for("home", submitted="1"))


@app.context_processor
def inject_submitted_flag():
    return {"submitted": request.args.get("submitted") == "1"}


if __name__ == "__main__":
    ensure_workbook()
    app.run(host="0.0.0.0", port=5000, debug=True)
